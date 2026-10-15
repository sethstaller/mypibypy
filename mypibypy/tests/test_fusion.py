"""Fusion-junction mode: gap on the breakpoint, partner-exclusive pads, both strands.

Two layers again. The synthetic layer runs everywhere. The reproduction layer
checks the tool against the lab's two hand-made fusion designs, and is skipped
unless a workbook is supplied out-of-band:

    MYPIBYPY_FUSION_WORKBOOK=/path/to/design.xlsx pytest tests/test_fusion.py

That workbook is gitignored and its sequences never enter this file, the
repository, or an assertion message -- the test reports positions and lengths,
never bases.

`TestArmGeometry` is the one to read. A fusion pair with its two pads on the
wrong arms is indistinguishable from a good one by length, composition, Tm or
any per-probe check; the only thing that reveals it is asking where the arms
point. That check found a real defect in 4 of 8 hand-made primaries.
"""
from __future__ import annotations

import os
import random
import re

import pytest

from mypibypy.design import (
    FOOTPRINT,
    GAP_FIVE_PRIME,
    GAP_THREE_PRIME,
    LEFT_LEN,
    PADS_AS_ORDERED,
    PADS_JUXTAPOSED,
    FusionGeometryError,
    arm_offsets,
    arms_are_juxtaposed,
    build_fusion_pairs,
    check_partner_specificity,
    fusion_candidate,
    fusion_footprint,
    pairs_from_footprint,
    split_footprint,
    verify_against,
)
from mypibypy.loci import FusionJunction, LociNotConfigured, load_fusions
from mypibypy.sequence import revcomp
from mypibypy.systems import systems


@pytest.fixture(scope="module")
def registry():
    reg = systems()
    assert reg, "no amplifier systems configured"
    return reg


@pytest.fixture(scope="module")
def two_systems(registry):
    vals = list(registry.values())
    return vals[0], (vals[1] if len(vals) > 1 else vals[0])


@pytest.fixture(scope="module")
def junction():
    """A synthetic junction whose two partners share no 25-mer."""
    rng = random.Random(99)
    five = "".join(rng.choice("ACGT") for _ in range(120))
    three = "".join(rng.choice("ACGT") for _ in range(120))
    return FusionJunction("demo", "GENEA", "GENEB", five + three, 120, "cy5", "cy3")


class TestGapSitsOnTheBreakpoint:
    """The point of fusion mode: the skipped base is the breakpoint itself."""

    def test_left_pad_is_all_five_prime_and_right_pad_all_three_prime(self):
        seq = "A" * 60 + "C" * 60
        for side in (GAP_FIVE_PRIME, GAP_THREE_PRIME):
            _, fp = fusion_footprint(seq, 60, side)
            left, gap, right = split_footprint(fp)
            assert set(left) == {"A"}, f"{side}: left pad must be 5' partner only"
            assert set(right) == {"C"}, f"{side}: right pad must be 3' partner only"
            assert len(left) == len(right) == 25 and len(fp) == FOOTPRINT

    def test_gap_side_decides_which_partner_gives_up_the_base(self):
        seq = "A" * 60 + "C" * 60
        _, five = fusion_footprint(seq, 60, GAP_FIVE_PRIME)
        _, three = fusion_footprint(seq, 60, GAP_THREE_PRIME)
        assert split_footprint(five)[1] == "A", "5' partner gives up its last base"
        assert split_footprint(three)[1] == "C", "3' partner gives up its first base"

    def test_the_two_gap_sides_differ_by_exactly_one_base_of_offset(self):
        seq = "".join(random.Random(3).choice("ACGT") for _ in range(200))
        a, _ = fusion_footprint(seq, 100, GAP_FIVE_PRIME)
        b, _ = fusion_footprint(seq, 100, GAP_THREE_PRIME)
        assert b - a == 1

    def test_too_little_partner_sequence_is_a_clear_error(self):
        with pytest.raises(FusionGeometryError) as exc:
            fusion_footprint("ACGT" * 10, 5, GAP_FIVE_PRIME)
        assert "5' partner" in str(exc.value)

    def test_footprint_must_be_51_nt(self):
        with pytest.raises(FusionGeometryError):
            split_footprint("ACGT" * 5)


class TestPartnerExclusivity:
    """Each pad must map to its own partner and never to the other.

    If the left pad is also present in the 3' partner, that partner alone
    presents both halves and the pair reports the unfused gene -- the exact thing
    a junction probe exists to rule out.
    """

    def test_a_clean_junction_has_no_problems(self, junction):
        _, fp = fusion_footprint(junction.sequence, junction.breakpoint)
        assert check_partner_specificity(
            fp, junction.five_prime_seq, junction.three_prime_seq) == []

    def test_left_pad_recurring_in_the_three_prime_partner_is_caught(self, junction):
        _, fp = fusion_footprint(junction.sequence, junction.breakpoint)
        left, _, _ = split_footprint(fp)
        contaminated = junction.three_prime_seq + left      # pad now in both partners
        problems = check_partner_specificity(fp, junction.five_prime_seq, contaminated)
        assert any("left pad also occurs in the 3' partner" in p for p in problems)
        assert any("not junction-specific" in p for p in problems)

    def test_right_pad_recurring_in_the_five_prime_partner_is_caught(self, junction):
        _, fp = fusion_footprint(junction.sequence, junction.breakpoint)
        _, _, right = split_footprint(fp)
        problems = check_partner_specificity(
            fp, junction.five_prime_seq + right, junction.three_prime_seq)
        assert any("right pad also occurs in the 5' partner" in p for p in problems)

    def test_reverse_complement_presence_counts_too(self, junction):
        """A pad matching the wrong partner's other strand recruits probe just as well."""
        _, fp = fusion_footprint(junction.sequence, junction.breakpoint)
        left, _, _ = split_footprint(fp)
        contaminated = junction.three_prime_seq + revcomp(left)
        problems = check_partner_specificity(fp, junction.five_prime_seq, contaminated)
        assert any("left pad also occurs" in p for p in problems)

    def test_a_pad_missing_from_its_own_partner_is_caught(self, junction):
        _, fp = fusion_footprint(junction.sequence, junction.breakpoint)
        problems = check_partner_specificity(fp, "ACGT" * 30, junction.three_prime_seq)
        assert any("not found in the 5' partner" in p for p in problems)


class TestBothStrandsAreEmitted:
    def test_a_junction_yields_a_sense_and_an_antisense_pair(self, junction, two_systems):
        sense_sys, anti_sys = two_systems
        pairs = build_fusion_pairs(junction.name, junction.sequence,
                                   junction.breakpoint, sense_sys, anti_sys)
        assert len(pairs) == 2
        assert "sense" in pairs[0].p1l_name and "antisense" in pairs[1].p1l_name
        assert pairs[0].system is sense_sys and pairs[1].system is anti_sys

    def test_the_two_pairs_go_on_different_channels(self, junction, two_systems):
        sense_sys, anti_sys = two_systems
        if sense_sys.name == anti_sys.name:
            pytest.skip("only one system configured")
        pairs = build_fusion_pairs(junction.name, junction.sequence,
                                   junction.breakpoint, sense_sys, anti_sys)
        assert pairs[0].system.fluorophore != pairs[1].system.fluorophore

    def test_antisense_footprint_is_the_reverse_complement_of_the_sense_one(self, junction):
        a = fusion_candidate(junction.name, junction.sequence, junction.breakpoint)
        b = fusion_candidate(junction.name, junction.sequence, junction.breakpoint,
                             antisense=True)
        assert b.left.seq == revcomp(a.right.seq)
        assert b.right.seq == revcomp(a.left.seq)
        assert b.gap_base == revcomp(a.gap_base)

    def test_antisense_pads_are_the_sense_pads_read_forward(self, junction, two_systems):
        """The identity worth knowing, and the reason the two pairs dimerize."""
        sense_sys, anti_sys = two_systems
        sense, anti = build_fusion_pairs(junction.name, junction.sequence,
                                         junction.breakpoint, sense_sys, anti_sys)
        sense_l_pad = sense.p1l[len(sense_sys.arm_left):]
        anti_r_pad = anti.p1r[: -len(anti_sys.arm_right)]
        assert anti_r_pad == revcomp(sense_l_pad)

    def test_the_strand_is_recorded_on_the_interval(self, junction):
        a = fusion_candidate(junction.name, junction.sequence, junction.breakpoint)
        b = fusion_candidate(junction.name, junction.sequence, junction.breakpoint,
                             antisense=True)
        assert "sense" in a.interval.build and "antisense" in b.interval.build
        assert a.interval.length == FOOTPRINT


class TestArmGeometry:
    """Do the arms meet at the gap? Nothing else about a pair reveals this."""

    def test_sense_pair_arms_meet_at_the_gap(self, junction, two_systems):
        sense_sys, _ = two_systems
        _, fp = fusion_footprint(junction.sequence, junction.breakpoint)
        sense, _ = build_fusion_pairs(junction.name, junction.sequence,
                                      junction.breakpoint, sense_sys)
        assert arm_offsets(sense, fp) == (LEFT_LEN - 1, LEFT_LEN + 1)
        assert arms_are_juxtaposed(sense, fp)

    def test_antisense_pair_arms_also_meet_at_the_gap(self, junction, two_systems):
        sense_sys, anti_sys = two_systems
        _, fp = fusion_footprint(junction.sequence, junction.breakpoint)
        _, anti = build_fusion_pairs(junction.name, junction.sequence,
                                     junction.breakpoint, sense_sys, anti_sys)
        assert arms_are_juxtaposed(anti, fp), (
            "the antisense pair must rebuild the P2 site too, or it makes no spot")
        assert arm_offsets(anti, fp) == (LEFT_LEN + 1, LEFT_LEN - 1)

    def test_swapped_pads_point_the_arms_away_from_the_junction(
            self, junction, two_systems):
        """The defect found in the hand-made designs, pinned down.

        `as_ordered` reproduces how the lab's two fusion antisense pairs were
        built. The pair looks well formed -- right lengths, right arms, right
        pads -- but both arms end up at the outer edges of the footprint instead
        of at the gap, so P2's landing site is never reconstituted.
        """
        sense_sys, anti_sys = two_systems
        _, fp = fusion_footprint(junction.sequence, junction.breakpoint)
        _, anti = build_fusion_pairs(junction.name, junction.sequence,
                                     junction.breakpoint, sense_sys, anti_sys,
                                     pad_assignment=PADS_AS_ORDERED)
        assert arm_offsets(anti, fp) == (0, FOOTPRINT - 1)
        assert not arms_are_juxtaposed(anti, fp)
        assert any("arms point away" in f for f in anti.candidate.flags)

    def test_the_two_conventions_differ_only_for_the_antisense_pair(
            self, junction, two_systems):
        sense_sys, anti_sys = two_systems
        a = build_fusion_pairs(junction.name, junction.sequence, junction.breakpoint,
                               sense_sys, anti_sys, pad_assignment=PADS_JUXTAPOSED)
        b = build_fusion_pairs(junction.name, junction.sequence, junction.breakpoint,
                               sense_sys, anti_sys, pad_assignment=PADS_AS_ORDERED)
        assert (a[0].p1l, a[0].p1r) == (b[0].p1l, b[0].p1r), "sense pair is unaffected"
        assert (a[1].p1l, a[1].p1r) != (b[1].p1l, b[1].p1r)

    def test_unknown_pad_assignment_is_rejected(self, junction, two_systems):
        with pytest.raises(ValueError, match="pad_assignment"):
            build_fusion_pairs(junction.name, junction.sequence, junction.breakpoint,
                               two_systems[0], pad_assignment="whatever")


class TestFromAFootprintDirectly:
    def test_a_supplied_footprint_reduces_to_the_same_geometry(
            self, junction, two_systems):
        sense_sys, anti_sys = two_systems
        _, fp = fusion_footprint(junction.sequence, junction.breakpoint)
        a = build_fusion_pairs(junction.name, junction.sequence, junction.breakpoint,
                               sense_sys, anti_sys)
        b = pairs_from_footprint(junction.name, fp, sense_sys, anti_sys)
        assert [(p.p1l, p.p1r) for p in a] == [(p.p1l, p.p1r) for p in b]

    def test_wrong_length_footprint_is_refused(self, two_systems):
        with pytest.raises(FusionGeometryError):
            pairs_from_footprint("x", "ACGT" * 10, two_systems[0])


class TestLociFile:
    def test_the_shipped_example_defines_a_synthetic_junction(self):
        import mypibypy.loci as loci_mod
        from pathlib import Path
        example = Path(loci_mod.__file__).parent / "loci.example.json"
        fusions = load_fusions(example)
        assert len(fusions) == 1
        f = fusions[0]
        assert f.five_prime_seq and f.three_prime_seq
        assert f.sequence[f.breakpoint :] == f.three_prime_seq
        assert f.antisense_system != f.system, "a junction is ordered on both channels"

    def test_a_file_with_no_fusions_is_not_an_error(self, tmp_path):
        p = tmp_path / "loci.json"
        p.write_text('{"loci": [{"name": "a", "gene": "G", "build": "mm10", '
                     '"chrom": "chr1", "start": 1, "end": 100, "system": "cy5"}]}')
        assert load_fusions(p) == []

    def test_a_breakpoint_outside_the_sequence_is_rejected(self, tmp_path):
        p = tmp_path / "loci.json"
        p.write_text('{"fusions": [{"name": "bad", "five_prime_gene": "A", '
                     '"three_prime_gene": "B", "sequence": "ACGTACGTAC", '
                     '"breakpoint": 999, "system": "cy5"}]}')
        with pytest.raises(LociNotConfigured, match="outside its"):
            load_fusions(p)

    def test_non_dna_sequence_is_rejected(self, tmp_path):
        p = tmp_path / "loci.json"
        p.write_text('{"fusions": [{"name": "bad", "five_prime_gene": "A", '
                     '"three_prime_gene": "B", "sequence": "ACGTXXGTAC", '
                     '"breakpoint": 5, "system": "cy5"}]}')
        with pytest.raises(LociNotConfigured, match="non-DNA"):
            load_fusions(p)

    def test_missing_fields_name_what_is_missing(self, tmp_path):
        p = tmp_path / "loci.json"
        p.write_text('{"fusions": [{"name": "bad", "sequence": "ACGT", '
                     '"breakpoint": 2, "system": "cy5"}]}')
        with pytest.raises(LociNotConfigured, match="five_prime_gene"):
            load_fusions(p)


# --- Reproduction of the lab's hand-made designs ---------------------------

WORKBOOK = os.environ.get("MYPIBYPY_FUSION_WORKBOOK")


def _read_designs(path):
    """Pull (footprint, ordered primaries) per junction sheet from a workbook.

    Deliberately structural: finds the 51-nt target and the primaries by shape,
    not by hardcoded cell references, and returns nothing else.
    """
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    designs = {}
    for sheet in wb.sheetnames:
        cells = {}
        for row in wb[sheet].iter_rows(values_only=True):
            c = [x for x in row if x is not None]
            if len(c) >= 2 and isinstance(c[0], str) and isinstance(c[1], str):
                v = re.sub(r"\s", "", c[1]).upper()
                if set(v) <= set("ACGTN") and len(v) >= 20:
                    cells[c[0].strip()] = v
        footprint = next((v for k, v in cells.items()
                          if len(v) == FOOTPRINT and "SENSE" in k.upper()
                          and "ANTISENSE" not in k.upper()), None)
        primaries = {k: v for k, v in cells.items()
                     if re.match(r"^pp[lr]_\d+$", k, re.I)}
        if footprint and len(primaries) == 4:
            designs[sheet] = (footprint, dict(sorted(
                primaries.items(), key=lambda kv: (kv[1] != primaries[min(primaries)],))))
    return designs


@pytest.fixture(scope="module")
def designs():
    d = _read_designs(WORKBOOK)
    if not d:
        pytest.skip(f"no fusion designs found in {WORKBOOK}")
    return d


@pytest.mark.skipif(not WORKBOOK, reason="set $MYPIBYPY_FUSION_WORKBOOK to run")
class TestReproducesHandMadeDesigns:
    """The acceptance criterion: rebuild the lab's two fusion designs.

    Sequences stay in the workbook. Assertions here are about counts and
    positions only.
    """

    def test_found_both_hand_made_junctions(self, designs):
        assert len(designs) >= 2, f"expected 2 junctions, found {sorted(designs)}"

    def test_as_ordered_convention_reproduces_the_primaries(self, designs, two_systems):
        """7 of 8 must rebuild byte-for-byte under the as-ordered convention."""
        sense_sys, anti_sys = two_systems
        matches = total = 0
        notes = []
        for sheet, (footprint, ordered) in designs.items():
            pairs = pairs_from_footprint(sheet, footprint, sense_sys, anti_sys,
                                         pad_assignment=PADS_AS_ORDERED)
            for row in verify_against(pairs, ordered):
                total += 1
                matches += row["match"]
                if not row["match"]:
                    notes.append(f"{sheet}/{row['ordered_name']}: {row.get('note')}")
        assert total == 8, f"expected 8 primaries, compared {total}"
        assert matches >= 7, f"only {matches}/8 reproduced: {notes}"

    def test_the_one_mismatch_is_a_single_base_truncation(self, designs, two_systems):
        """Not a rule difference: one ordered primary is 1 nt short."""
        sense_sys, anti_sys = two_systems
        off = []
        for sheet, (footprint, ordered) in designs.items():
            pairs = pairs_from_footprint(sheet, footprint, sense_sys, anti_sys,
                                         pad_assignment=PADS_AS_ORDERED)
            off += [r for r in verify_against(pairs, ordered) if not r["match"]]
        assert len(off) <= 1
        for row in off:
            assert row["built_len"] - row["ordered_len"] == 1, row

    def test_the_ordered_antisense_pairs_are_not_juxtaposed(self, designs, two_systems):
        """The defect. Every antisense pair as ordered has its arms on the outside."""
        sense_sys, anti_sys = two_systems
        for sheet, (footprint, _) in designs.items():
            sense, anti = pairs_from_footprint(sheet, footprint, sense_sys, anti_sys,
                                               pad_assignment=PADS_AS_ORDERED)
            assert arms_are_juxtaposed(sense, footprint), f"{sheet}: sense pair is fine"
            assert not arms_are_juxtaposed(anti, footprint), (
                f"{sheet}: antisense pair unexpectedly juxtaposed -- the geometric "
                f"model in design.arm_offsets may be wrong, re-check before acting")
            assert arm_offsets(anti, footprint) == (0, FOOTPRINT - 1)

    def test_the_corrected_convention_fixes_the_geometry(self, designs, two_systems):
        sense_sys, anti_sys = two_systems
        for sheet, (footprint, _) in designs.items():
            for pair in pairs_from_footprint(sheet, footprint, sense_sys, anti_sys,
                                             pad_assignment=PADS_JUXTAPOSED):
                assert arms_are_juxtaposed(pair, footprint), f"{sheet}/{pair.p1l_name}"
