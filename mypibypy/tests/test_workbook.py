"""The order workbook must never claim a screen that did not run.

This is a regression test for a real incident: an earlier workbook asserted a
specificity result unconditionally, and an order sheet went out carrying an
unverified "repeat-free single-copy" claim. The guard is
`SpecificityReport.screened`, and these tests hold the workbook to it.

Nothing here touches the network: the region is synthetic and the candidates are
scanned locally.
"""
from __future__ import annotations

import random

import pytest
from openpyxl import load_workbook

from mypibypy.design import build_pair, scan
from mypibypy.loci import Locus
from mypibypy.pipeline import DesignResult
from mypibypy.sequence import GenomicInterval
from mypibypy.specificity import (
    SCREENED,
    Hit,
    SpecificityReport,
    not_screened,
    pair_hits,
)
from mypibypy.systems import systems
from mypibypy.workbook import write_workbook


@pytest.fixture(scope="module")
def sys_a():
    reg = systems()
    assert reg, "no amplifier systems configured"
    return list(reg.values())[0]


@pytest.fixture(scope="module")
def region():
    """A synthetic 300-nt region with balanced composition, scanned locally."""
    rng = random.Random(7)
    seq = "".join(rng.choice("ACGT") for _ in range(300))
    iv = GenomicInterval("mm10", "chr1", 1_000_000, 1_000_299)
    return seq, iv


def make_result(sys_a, region, spec: SpecificityReport) -> DesignResult:
    seq, iv = region
    locus = Locus("testlocus", "Gene1", "synthetic test region", iv, sys_a.name)
    cands = scan(seq, iv, top=3)
    assert cands
    pair = build_pair(locus.name, sys_a, cands[0], "P1L_test", "P1R_test")
    return DesignResult(locus, sys_a, seq, cands, pair, rmsk=[], specificity=spec)


def workbook_text(tmp_path, results) -> str:
    out = tmp_path / "order.xlsx"
    write_workbook(str(out), results, "unit test")
    wb = load_workbook(out)
    chunks = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            chunks.extend(str(v) for v in row if v is not None)
    return "\n".join(chunks)


class TestUnscreenedIsStated:
    def test_not_screened_says_so_in_capitals(self, tmp_path, sys_a, region):
        res = make_result(sys_a, region, not_screened("no local database configured"))
        text = workbook_text(tmp_path, [res])
        assert "NO GENOME-WIDE SPECIFICITY SCREEN WAS RUN" in text
        assert "no local database configured" in text

    def test_not_screened_never_claims_single_copy(self, tmp_path, sys_a, region):
        res = make_result(sys_a, region, not_screened("no database"))
        text = workbook_text(tmp_path, [res]).lower()
        assert "single-copy" in text, "the claim should be discussed..."
        # ...but only ever to deny it.
        for line in text.splitlines():
            if "single-copy" in line:
                assert ("not establish" in line or "nothing here establishes" in line), line

    def test_skipped_screen_reads_differently_from_a_clean_one(
            self, tmp_path, sys_a, region):
        """'We did not look' and 'we looked and it was clean' must not match."""
        skipped = workbook_text(
            tmp_path, [make_result(sys_a, region, not_screened("skipped by request"))])
        clean = workbook_text(tmp_path, [make_result(sys_a, region, SpecificityReport(
            status=SCREENED, backend="bowtie v1 -v 3 -a", database="mm10"))])
        assert "NO GENOME-WIDE SPECIFICITY SCREEN WAS RUN" in skipped
        assert "NO GENOME-WIDE SPECIFICITY SCREEN WAS RUN" not in clean


class TestScreenedIsReported:
    def test_clean_screen_reports_pass_and_the_backend(self, tmp_path, sys_a, region):
        rep = SpecificityReport(status=SCREENED, backend="bowtie v1 -v 3 -a (exhaustive)",
                                database="/data/mm10")
        text = workbook_text(tmp_path, [make_result(sys_a, region, rep)])
        assert "PASS" in text
        assert "bowtie v1 -v 3 -a" in text, "the method must be on the sheet"

    def test_a_real_offtarget_pair_is_reported_as_fail(self, tmp_path, sys_a, region):
        left = Hit("left", "chr9", 500, 524, "+", 0, 0, 25, 25, 1, 25)
        right = Hit("right", "chr9", 526, 550, "+", 0, 0, 25, 25, 1, 25)
        seq, _ = region
        cands = scan(seq, region[1], top=1)
        sites = pair_hits([left], [right], cands[0].left.seq, cands[0].right.seq)
        rep = SpecificityReport(status=SCREENED, backend="bowtie v1 -v 3 -a",
                                database="/data/mm10", left_seq=cands[0].left.seq,
                                right_seq=cands[0].right.seq, left_hits=[left],
                                right_hits=[right], pair_sites=sites)
        assert rep.verdict == "fail"
        text = workbook_text(tmp_path, [make_result(sys_a, region, rep)])
        assert "FAIL" in text and "do not order" in text
        assert "chr9" in text, "the off-target site itself must be on the sheet"

    def test_single_half_hits_do_not_read_as_failure(self, tmp_path, sys_a, region):
        """A lone-half hit cannot make a spot, so it must not read like one."""
        lone = Hit("left", "chr9", 500, 524, "+", 0, 0, 25, 25, 1, 25)
        rep = SpecificityReport(status=SCREENED, backend="bowtie v1 -v 3 -a",
                                database="/data/mm10", left_hits=[lone])
        assert rep.verdict != "fail"
        text = workbook_text(tmp_path, [make_result(sys_a, region, rep)])
        assert "do not order" not in text
