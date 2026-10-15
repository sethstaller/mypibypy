"""Run the cross-hybridization checks over the validated designs.

Point it at a directory of piFISH order workbooks. Those workbooks contain probe
and target sequences, so they are NOT committed to this repository -- keep them
locally (e.g. in a gitignored designs/ directory) and pass that path:

    python scripts/calibrate_crosshyb.py ~/path/to/designs

Orders that were synthesized and imaged successfully are the calibration set: any
FLAG they raise is a false positive and the threshold needs loosening. That is the
whole point of running them.

Requires configured amplifier systems (systems.local.json); the synthetic example
systems will not match real workbooks' arms.
"""
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook

from mypibypy.crosshyb import check_order, check_pair
from mypibypy.systems import is_example, loaded_from, systems


def system_of(seq: str):
    """Identify a primary's system from its arm sequence, not its name."""
    for s in systems().values():
        if seq.upper().startswith(s.arm_left) or seq.upper().endswith(s.arm_right):
            return s
    return None


def primaries_from(path: Path) -> dict[str, str]:
    wb = load_workbook(path, data_only=True)
    out = {}
    for sheet in ("ORDER_ssDNA", "Parts list"):
        if sheet not in wb.sheetnames:
            continue
        for row in wb[sheet].iter_rows(values_only=True):
            if not row or not row[0] or len(row) < 2 or not isinstance(row[1], str):
                continue
            name, seq = str(row[0]), row[1].strip().upper()
            if name.upper().startswith(("P1L", "P1R")) and set(seq) <= set("ACGT"):
                out.setdefault(name, seq)
    return out


def main() -> None:
    root = Path(sys.argv[1] if len(sys.argv) > 1 else ".")
    books = sorted(root.glob("*piFISH*.xlsx"))
    if not books:
        raise SystemExit(f"no piFISH workbooks found under {root.resolve()}")
    systems()  # load before reporting where from
    print(f"systems loaded from {loaded_from()}")
    if is_example():
        print("WARNING: example (synthetic) systems are loaded -- arms will not "
              "match real workbooks, so every pair will be skipped.\n")
    all_primaries = {}
    for book in books:
        prim = primaries_from(book)
        print(f"\n=== {book.name}: {len(prim)} ordered primaries")
        pairs = {}
        for name, seq in prim.items():
            key = name.split("_", 1)[1] if "_" in name else name
            pairs.setdefault(key, {})[name[:3].upper()] = (name, seq)
        for key, sides in pairs.items():
            if "P1L" not in sides or "P1R" not in sides:
                print(f"  [skip] {key}: incomplete pair {list(sides)}")
                continue
            (ln, lseq), (rn, rseq) = sides["P1L"], sides["P1R"]
            sysn = system_of(lseq) or system_of(rseq)
            if sysn is None:
                print(f"  [skip] {key}: no known system matches {ln}")
                continue
            rep = check_pair(lseq, rseq, sysn, p1l_name=ln, p1r_name=rn)
            print(f"  -- {key}  [{sysn.name}/{sysn.fluorophore}]  {rep.summary()}")
            for f in sorted(rep.findings, key=lambda f: -f.duplex.length):
                if f.flagged or f.duplex.length >= 10:
                    print(f"     {f}")
            all_primaries.update({ln: (lseq, sysn), rn: (rseq, sysn)})

    print(f"\n=== whole-pool cross-pairing across all {len(all_primaries)} primaries")
    rep = check_order(all_primaries)
    print(f"  {rep.summary()}")
    for f in sorted(rep.findings, key=lambda f: -f.duplex.length)[:12]:
        print(f"     {f}")


if __name__ == "__main__":
    main()
